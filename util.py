"""Shared helpers: raw sheet reading, month parsing, name normalisation."""
from __future__ import annotations

import datetime as dt
import re

import pandas as pd

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def grid(path: str, sheet: str, max_col: int = 80, max_row: int | None = None) -> list[list]:
    """Raw cell values of one sheet (xlsx via openpyxl read-only, xlsb via pyxlsb), trailing empties trimmed."""
    rows: list[list] = []
    if path.lower().endswith(".xlsb"):
        from pyxlsb import open_workbook
        with open_workbook(path) as wb:
            with wb.get_sheet(sheet) as sh:
                for i, r in enumerate(sh.rows()):
                    if max_row and i >= max_row: break
                    rows.append([c.v for c in r][:max_col])
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet]
        empty = 0
        for r in ws.iter_rows(max_col=max_col, max_row=max_row, values_only=True):
            r = list(r)
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
                empty += 1
                if empty > 500: break
            else:
                empty = 0
            rows.append(r)
        wb.close()
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    w = max((len(r) for r in rows), default=0)
    return [list(r) + [None] * (w - len(r)) for r in rows]


def sheet_names(path: str) -> list[str]:
    if path.lower().endswith(".xlsb"):
        from pyxlsb import open_workbook
        with open_workbook(path) as wb:
            return list(wb.sheets)
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    n = wb.sheetnames; wb.close(); return n


def find_row(g: list[list], pred, limit: int = 20) -> int:
    for i, r in enumerate(g[:limit]):
        if pred(r): return i
    return -1


def ffill(row: list) -> list:
    out, last = [], None
    for v in row:
        if v is not None and str(v).strip() != "": last = v
        out.append(last)
    return out


def s(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def num(v) -> float:
    if v is None: return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return 0.0 if pd.isna(v) else float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return 0.0


def code(v) -> str:
    """Dealer codes arrive as 1348, 1348.0, '1348 ', 'AU30006' - normalise to a clean string."""
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return re.sub(r"\s+", "", str(v)).strip().upper()


def person(v) -> str:
    """Upper-case, single-spaced person name with known aliases folded."""
    n = s(v).upper()
    return ALIASES.get(n, n)


ALIASES = {
    "VIKAS GUPTA": "VIKAS G",
    "CHANDAN CHATERJEE": "CHANDAN CHATTERJEE",
    "BISWABHANU": "BISWABHANU BISWAL",
    "ROHIT GANDHI": "ROHITH GANDHI",
    "LOKNATH PATTAJOSHI": "LOKANATH PATTAJOSHI",
}


def month_key(label, default_year: int | None = None) -> str | None:
    """"Aug'26" / "Aug`26" / "Aug'`26" / "August'26" / "Jan'26 MONTH" / "Jul-26" / datetime -> '2026-08'."""
    if isinstance(label, (dt.datetime, dt.date, pd.Timestamp)):
        return f"{label.year:04d}-{label.month:02d}"
    t = s(label).lower()
    m = re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s'`’\-]*`?'?(\d{2,4})\b", t)
    if m:
        y = int(m.group(2)); y = y + 2000 if y < 100 else y
        return f"{y:04d}-{MONTHS[m.group(1)]:02d}"
    m = re.fullmatch(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*", t)
    if m and default_year:
        return f"{default_year:04d}-{MONTHS[m.group(1)]:02d}"
    return None


def month_label(key: str) -> str:
    y, m = key.split("-")
    return dt.date(int(y), int(m), 1).strftime("%b'%y")


def date_from_name(name: str) -> dt.date | None:
    """As-of date written into MIS file names: 28-Sep-2026, 28-Sep-26, 28-Sep'26, 21 Sep 2026, 01082026."""
    t = name.replace("’", "'")
    m = re.search(r"(\d{1,2})(?:st|nd|rd|th)?[\s\-_']*(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s\-_'`]*(\d{2,4})", t, re.I)
    if m:
        y = int(m.group(3)); y = y + 2000 if y < 100 else y
        try: return dt.date(y, MONTHS[m.group(2).lower()], int(m.group(1)))
        except ValueError: return None
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", t)
    if m:
        try: return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError: return None
    return None


def ratio(a, b) -> float:
    a, b = float(a or 0), float(b or 0)
    return a / b if b else float("nan")


def excel_date(v):
    if isinstance(v, (dt.datetime, dt.date)): return pd.Timestamp(v)
    try:
        f = float(v)
        if 20000 < f < 80000: return pd.Timestamp("1899-12-30") + pd.Timedelta(days=f)
    except (TypeError, ValueError):
        pass
    return pd.NaT


def drop_total_rows(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    """Remove hidden grand-total rows: a row whose value equals the sum of all others in a numeric column."""
    d = df.copy()
    for c in cols:
        if c not in d or d.empty: continue
        v = pd.to_numeric(d[c], errors="coerce").fillna(0)
        tot = v.sum(); mx = v.max()
        if mx > 0 and abs((tot - mx) - mx) < 1e-6 * max(1, mx):
            d = d.drop(index=v.idxmax())
    return d

"""Report model + exporters (HTML, PDF, PPTX, CSV).

Every screen, dashboard and chat answer is built as a `Report` (a list of blocks), so what
the user sees is exactly what gets exported.
"""
from __future__ import annotations

import datetime as dt
import html
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

import pandas as pd

from .charts import ChartSpec, to_plotly, to_png

BRAND = "#1F3A5F"
ACCENT = "#E4572E"


# --------------------------------------------------------------------------- #
# model
# --------------------------------------------------------------------------- #
@dataclass
class Block:
    kind: str                     # heading | text | bullets | kpis | table | chart
    title: str = ""
    text: str = ""
    items: list = field(default_factory=list)     # bullets or kpis [(label, value, note)]
    df: pd.DataFrame | None = None
    chart: ChartSpec | None = None


@dataclass
class Report:
    title: str
    subtitle: str = ""
    blocks: list[Block] = field(default_factory=list)

    # builder helpers -------------------------------------------------------
    def h(self, t):  self.blocks.append(Block("heading", title=t)); return self
    def p(self, t):  self.blocks.append(Block("text", text=t)); return self
    def bullets(self, items, title=""):
        if items: self.blocks.append(Block("bullets", title=title, items=list(items)))
        return self
    def kpis(self, items, title=""):
        self.blocks.append(Block("kpis", title=title, items=list(items))); return self
    def table(self, df, title=""):
        if df is not None and len(df): self.blocks.append(Block("table", title=title, df=df))
        return self
    def chart(self, spec: ChartSpec):
        if spec is not None: self.blocks.append(Block("chart", title=spec.title, chart=spec))
        return self
    def extend(self, other: "Report"):
        self.blocks.extend(other.blocks); return self

    def slug(self):
        return re.sub(r"[^A-Za-z0-9]+", "_", self.title).strip("_")[:60] or "report"


# --------------------------------------------------------------------------- #
# formatting
# --------------------------------------------------------------------------- #
PCT_HINT = re.compile(r"(%|pct|ratio|share|conv|retention|penetration|rate)", re.I)


def fmt_cell(col: str, v: Any) -> str:
    try:
        if v is None or pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, (bool,)) or type(v).__name__ == "bool_":
        return "Yes" if v else "No"
    if isinstance(v, (pd.Timestamp, dt.date, dt.datetime)):
        return pd.Timestamp(v).strftime("%d-%b-%Y")
    if isinstance(v, float) or isinstance(v, int):
        if PCT_HINT.search(str(col)) and isinstance(v, float) and -5 <= v <= 5:
            return f"{v * 100:.1f}%"
        if isinstance(v, float) and abs(v - round(v)) > 1e-9:
            return f"{v:,.2f}"
        return f"{int(round(v)):,}"
    return str(v)


def display_df(df: pd.DataFrame, max_rows: int | None = None) -> pd.DataFrame:
    d = df if max_rows is None else df.head(max_rows)
    return pd.DataFrame({c: [fmt_cell(c, v) for v in d[c]] for c in d.columns})


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
CSS = """
:root{--brand:%s;--accent:%s;--ink:#1d2433;--muted:#667085;--line:#e4e7ec;--bg:#fff;--tile:#f6f8fb}
@media (prefers-color-scheme: dark){:root{--ink:#e7eaf0;--muted:#98a2b3;--line:#2c3444;--bg:#111722;--tile:#18202e}}
*{box-sizing:border-box}body{font-family:Inter,Segoe UI,Arial,sans-serif;color:var(--ink);background:var(--bg);margin:0;padding:24px 16px;line-height:1.45}
main{max-width:1100px;margin:auto}h1{color:var(--brand);margin:0 0 4px;font-size:26px}h2{border-bottom:2px solid var(--brand);padding-bottom:4px;margin-top:28px;font-size:19px}
h3{font-size:15px;margin:18px 0 6px}.sub{color:var(--muted);margin-bottom:18px}
.kpis{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:10px;margin:10px 0}
.kpi{background:var(--tile);border:1px solid var(--line);border-radius:10px;padding:10px 12px}.kpi .l{font-size:12px;color:var(--muted)}.kpi .v{font-size:22px;font-weight:700}.kpi .n{font-size:12px;color:var(--muted)}
.tw{overflow-x:auto;margin:8px 0 14px}table{border-collapse:collapse;font-size:12.5px;width:100%%}th{background:var(--brand);color:#fff;text-align:left;padding:6px 8px;white-space:nowrap}
td{padding:5px 8px;border-bottom:1px solid var(--line);white-space:nowrap}tr:nth-child(even) td{background:var(--tile)}
ul{margin:6px 0 12px}footer{color:var(--muted);font-size:11px;margin-top:30px}
@media print{body{padding:0}.chart{page-break-inside:avoid}}
""" % (BRAND, ACCENT)


def to_html(rep: Report) -> bytes:
    out = [f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>{html.escape(rep.title)}</title><style>{CSS}</style>"
           "<script src='https://cdn.jsdelivr.net/npm/plotly.js-dist-min@2.35.2/plotly.min.js'></script></head><body><main>",
           f"<h1>{html.escape(rep.title)}</h1><div class='sub'>{html.escape(rep.subtitle)}</div>"]
    for b in rep.blocks:
        if b.kind == "heading":
            out.append(f"<h2>{html.escape(b.title)}</h2>")
        elif b.kind == "text":
            out.append("".join(f"<p>{_md_inline(x)}</p>" for x in b.text.split("\n\n") if x.strip()))
        elif b.kind == "bullets":
            if b.title: out.append(f"<h3>{html.escape(b.title)}</h3>")
            out.append("<ul>" + "".join(f"<li>{_md_inline(str(i))}</li>" for i in b.items) + "</ul>")
        elif b.kind == "kpis":
            if b.title: out.append(f"<h3>{html.escape(b.title)}</h3>")
            out.append("<div class='kpis'>" + "".join(
                f"<div class='kpi'><div class='l'>{html.escape(str(k[0]))}</div><div class='v'>{html.escape(str(k[1]))}</div>"
                f"<div class='n'>{html.escape(str(k[2]) if len(k) > 2 and k[2] else '')}</div></div>" for k in b.items) + "</div>")
        elif b.kind == "table":
            if b.title: out.append(f"<h3>{html.escape(b.title)}</h3>")
            out.append("<div class='tw'>" + display_df(b.df).to_html(index=False, escape=True, border=0) + "</div>")
        elif b.kind == "chart":
            fig = to_plotly(b.chart)
            out.append("<div class='chart'>" + fig.to_html(full_html=False, include_plotlyjs=False,
                                                          config={"displaylogo": False, "responsive": True}) + "</div>")
    out.append(f"<footer>Generated {dt.datetime.now():%d %b %Y %H:%M} · MIS Agent</footer></main></body></html>")
    return "".join(out).encode("utf-8")


def _md_inline(s: str) -> str:
    s = html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    return s


# --------------------------------------------------------------------------- #
# PDF (fpdf2)
# --------------------------------------------------------------------------- #
def _latin(s: str) -> str:
    s = str(s).replace("₹", "Rs ").replace("–", "-").replace("—", "-").replace("’", "'").replace("…", "...") \
        .replace("→", "->").replace("≥", ">=").replace("≤", "<=").replace("Δ", "chg ").replace("•", "-").replace("›", ">")
    return s.encode("latin-1", "replace").decode("latin-1")


def to_pdf(rep: Report) -> bytes:
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_fill_color(31, 58, 95); self.rect(0, 0, self.w, 9, "F")
            self.set_y(2); self.set_font("Helvetica", "B", 8); self.set_text_color(255, 255, 255)
            self.cell(0, 5, _latin(rep.title), align="L")
            self.set_text_color(29, 36, 51); self.set_y(13)

        def footer(self):
            self.set_y(-10); self.set_font("Helvetica", "", 7); self.set_text_color(120, 120, 120)
            self.cell(0, 5, f"Generated {dt.datetime.now():%d %b %Y %H:%M}  |  Page {self.page_no()}", align="R")

    pdf = PDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(True, 12)
    pdf.add_page()
    W = pdf.w - 20
    pdf.set_font("Helvetica", "B", 20); pdf.set_text_color(31, 58, 95)
    pdf.multi_cell(W, 9, _latin(rep.title), new_x="LMARGIN", new_y="NEXT")
    if rep.subtitle:
        pdf.set_font("Helvetica", "", 10); pdf.set_text_color(100, 100, 100); pdf.multi_cell(W, 5, _latin(rep.subtitle), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2); pdf.set_text_color(29, 36, 51)

    for b in rep.blocks:
        if b.kind == "heading":
            if pdf.get_y() > pdf.h - 70: pdf.add_page()
            pdf.ln(2); pdf.set_font("Helvetica", "B", 13); pdf.set_text_color(31, 58, 95)
            pdf.cell(W, 7, _latin(b.title), new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(31, 58, 95); pdf.line(10, pdf.get_y(), 10 + W, pdf.get_y()); pdf.ln(2)
            pdf.set_text_color(29, 36, 51)
        elif b.kind == "text":
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(W, 5, _latin(b.text.replace("**", "")), new_x="LMARGIN", new_y="NEXT"); pdf.ln(1)
        elif b.kind == "bullets":
            if b.title: _sub(pdf, b.title, W)
            pdf.set_font("Helvetica", "", 10)
            for i in b.items:
                pdf.multi_cell(W, 5, _latin("- " + str(i).replace("**", "")), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif b.kind == "kpis":
            if b.title: _sub(pdf, b.title, W)
            n = min(5, max(1, len(b.items))); cw = W / n
            for r0 in range(0, len(b.items), n):
                if pdf.get_y() > pdf.h - 30: pdf.add_page()
                y = pdf.get_y()
                for j, k in enumerate(b.items[r0:r0 + n]):
                    x = 10 + j * cw
                    pdf.set_fill_color(246, 248, 251); pdf.set_draw_color(220, 224, 230)
                    pdf.rect(x + 1, y, cw - 2, 18, "DF")
                    pdf.set_xy(x + 3, y + 1.5); pdf.set_font("Helvetica", "", 7.5); pdf.set_text_color(100, 100, 100)
                    pdf.cell(cw - 6, 4, _latin(k[0])[:45])
                    pdf.set_xy(x + 3, y + 6); pdf.set_font("Helvetica", "B", 13); pdf.set_text_color(29, 36, 51)
                    pdf.cell(cw - 6, 6, _latin(k[1])[:22])
                    if len(k) > 2 and k[2]:
                        pdf.set_xy(x + 3, y + 12.5); pdf.set_font("Helvetica", "", 7); pdf.set_text_color(100, 100, 100)
                        pdf.cell(cw - 6, 4, _latin(k[2])[:50])
                pdf.set_y(y + 21)
            pdf.set_text_color(29, 36, 51)
        elif b.kind == "table":
            if b.title: _sub(pdf, b.title, W)
            _pdf_table(pdf, display_df(b.df), W)
        elif b.kind == "chart":
            png = to_png(b.chart, w=10.5, h=4.0)
            cw = min(W, 205); ch = cw * 4.0 / 10.5
            if pdf.get_y() + ch > pdf.h - 12: pdf.add_page()
            pdf.image(io.BytesIO(png), x=10, w=cw); pdf.ln(3)
    return bytes(pdf.output())


def _sub(pdf, t, W):
    if pdf.get_y() > pdf.h - 30: pdf.add_page()
    pdf.set_font("Helvetica", "B", 10.5); pdf.cell(W, 6, _latin(t), new_x="LMARGIN", new_y="NEXT")


def _pdf_table(pdf, d: pd.DataFrame, W, max_rows=400):
    cols = list(d.columns)
    if not cols: return
    d = d.head(max_rows)
    # width by content, capped
    lens = [max([len(str(c))] + [len(str(v)) for v in d[c].head(60)]) for c in cols]
    lens = [min(max(l, 4), 34) for l in lens]
    fs = 8 if len(cols) <= 10 else 7 if len(cols) <= 14 else 6
    tot = sum(lens); widths = [W * l / tot for l in lens]
    maxch = [max(3, int(w / (fs * 0.2))) for w in widths]

    def row(vals, head=False):
        if pdf.get_y() > pdf.h - 16:
            pdf.add_page(); row(cols, True)
        pdf.set_font("Helvetica", "B" if head else "", fs)
        if head: pdf.set_fill_color(31, 58, 95); pdf.set_text_color(255, 255, 255)
        for w, v, m in zip(widths, vals, maxch):
            s = _latin(v); s = s if len(s) <= m else s[:m - 1] + "."
            pdf.cell(w, 5, s, border=0, fill=head)
        pdf.ln(5)
        if head: pdf.set_text_color(29, 36, 51)
        else:
            pdf.set_draw_color(228, 231, 236); pdf.line(10, pdf.get_y(), 10 + W, pdf.get_y())

    row(cols, True)
    for _, r in d.iterrows(): row([r[c] for c in cols])
    pdf.ln(3)


# --------------------------------------------------------------------------- #
# PPTX (python-pptx, native editable charts + tables)
# --------------------------------------------------------------------------- #
def to_pptx(rep: Report) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    navy = RGBColor(0x1F, 0x3A, 0x5F)

    def new_slide(title):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.9)); bar.fill.solid(); bar.fill.fore_color.rgb = navy
        bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6)).text_frame
        tb.text = title[:90]; r = tb.paragraphs[0].runs[0]; r.font.size = Pt(24); r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        return s

    def text_box(s, lines, top=1.1, size=14):
        tf = s.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(6.2 - (top - 1.1))).text_frame
        tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = str(ln).replace("**", ""); p.font.size = Pt(size); p.space_after = Pt(6)

    # title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = navy
    bg.line.fill.background()
    tf = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(2)).text_frame; tf.word_wrap = True
    tf.text = rep.title; tf.paragraphs[0].runs[0].font.size = Pt(38); tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p = tf.add_paragraph(); p.text = rep.subtitle or f"Generated {dt.datetime.now():%d %b %Y}"
    p.font.size = Pt(16); p.font.color.rgb = RGBColor(0xD0, 0xD8, 0xE4)

    section = rep.title
    pending_text: list[str] = []

    def flush_text():
        nonlocal pending_text
        if pending_text:
            text_box(new_slide(section), pending_text[:14], size=15 if len(pending_text) < 9 else 12)
            pending_text = []

    for b in rep.blocks:
        if b.kind == "heading":
            flush_text(); section = b.title
        elif b.kind == "text":
            pending_text += [x for x in b.text.split("\n") if x.strip()]
        elif b.kind == "bullets":
            if b.title: pending_text.append(b.title + ":")
            pending_text += ["• " + str(i) for i in b.items]
        elif b.kind == "kpis":
            flush_text()
            sl = new_slide(b.title or section)
            items = b.items[:12]; per = 4; w = 12.3 / per
            for i, k in enumerate(items):
                x = 0.5 + (i % per) * w; y = 1.2 + (i // per) * 1.8
                box = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w - 0.2), Inches(1.6))
                box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF3, 0xF5, 0xF9); box.line.color.rgb = RGBColor(0xDD, 0xE2, 0xEA)
                t = box.text_frame; t.clear(); t.word_wrap = True
                t.paragraphs[0].text = str(k[0]); t.paragraphs[0].font.size = Pt(12); t.paragraphs[0].font.color.rgb = RGBColor(0x66, 0x70, 0x85)
                pv = t.add_paragraph(); pv.text = str(k[1]); pv.font.size = Pt(26); pv.font.bold = True; pv.font.color.rgb = navy
                if len(k) > 2 and k[2]:
                    pn = t.add_paragraph(); pn.text = str(k[2]); pn.font.size = Pt(10); pn.font.color.rgb = RGBColor(0x66, 0x70, 0x85)
        elif b.kind == "table":
            flush_text()
            d = display_df(b.df)
            cols = list(d.columns)[:12]
            for start in range(0, max(1, min(len(d), 60)), 14):
                chunk = d.iloc[start:start + 14]
                sl = new_slide((b.title or section) + (f" ({start // 14 + 1})" if len(d) > 14 else ""))
                shp = sl.shapes.add_table(len(chunk) + 1, len(cols), Inches(0.4), Inches(1.1), Inches(12.5),
                                          Inches(0.35 * (len(chunk) + 1)))
                tbl = shp.table
                fs = Pt(11 if len(cols) <= 7 else 9 if len(cols) <= 10 else 8)
                for j, c in enumerate(cols):
                    cell = tbl.cell(0, j); cell.text = str(c); cell.text_frame.paragraphs[0].font.size = fs
                    cell.text_frame.paragraphs[0].font.bold = True
                for i, (_, r) in enumerate(chunk.iterrows(), 1):
                    for j, c in enumerate(cols):
                        cell = tbl.cell(i, j); cell.text = str(r[c])[:40]; cell.text_frame.paragraphs[0].font.size = fs
        elif b.kind == "chart":
            flush_text()
            spec = b.chart
            sl = new_slide(spec.title or section)
            if spec.kind in ("bar", "hbar", "line", "stacked", "pie"):
                cd = CategoryChartData(); cd.categories = [str(x)[:30] for x in spec.x]
                for name, ys in spec.series.items():
                    cd.add_series(name, [None if (y is None or pd.isna(y)) else float(y) for y in ys])
                kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
                        "line": XL_CHART_TYPE.LINE_MARKERS, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
                        "pie": XL_CHART_TYPE.PIE, "combo": XL_CHART_TYPE.COLUMN_CLUSTERED}[spec.kind]
                gf = sl.shapes.add_chart(kind, Inches(0.5), Inches(1.1), Inches(12.3), Inches(6.1), cd)
                ch = gf.chart
                ch.has_legend = len(spec.series) > 1 or spec.kind == "pie"
                if ch.has_legend:
                    ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
                if spec.pct and spec.kind != "pie":
                    ch.value_axis.tick_labels.number_format = '0%'; ch.value_axis.tick_labels.number_format_is_linked = False
                if spec.kind != "pie":
                    pl = ch.plots[0]; pl.has_data_labels = len(spec.x) <= 15
                    if pl.has_data_labels:
                        pl.data_labels.font.size = Pt(9)
                        pl.data_labels.number_format = '0.0%' if spec.pct else '#,##0'
                        pl.data_labels.number_format_is_linked = False
            else:
                sl.shapes.add_picture(io.BytesIO(to_png(spec, w=12, h=5.6)), Inches(0.5), Inches(1.1), width=Inches(12.3))
    flush_text()
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def to_csv(rep: Report) -> tuple[bytes, str]:
    """One table -> .csv; several -> .zip of CSVs (charts exported as their data)."""
    tables = []
    for b in rep.blocks:
        if b.kind == "table":
            tables.append((b.title or "table", b.df))
        elif b.kind == "chart":
            tables.append(((b.chart.title or "chart") + " (chart data)", b.chart.to_df()))
        elif b.kind == "kpis":
            tables.append((b.title or "KPIs", pd.DataFrame([{"Metric": k[0], "Value": k[1], "Note": k[2] if len(k) > 2 else ""} for k in b.items])))
    if not tables:
        return b"", "csv"
    if len(tables) == 1:
        return tables[0][1].to_csv(index=False).encode("utf-8-sig"), "csv"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        seen = {}
        for name, df in tables:
            n = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")[:50] or "table"
            seen[n] = seen.get(n, 0) + 1
            if seen[n] > 1: n += f"_{seen[n]}"
            z.writestr(f"{n}.csv", df.to_csv(index=False).encode("utf-8-sig"))
    return buf.getvalue(), "zip"


def export(rep: Report, fmt: str) -> tuple[bytes, str, str]:
    """-> (bytes, filename, mime)."""
    stamp = dt.date.today().strftime("%Y-%m-%d")
    base = f"{rep.slug()}_{stamp}"
    if fmt == "html":
        return to_html(rep), base + ".html", "text/html"
    if fmt == "pdf":
        return to_pdf(rep), base + ".pdf", "application/pdf"
    if fmt == "pptx":
        return to_pptx(rep), base + ".pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    data, ext = to_csv(rep)
    return data, f"{base}.{ext}", "text/csv" if ext == "csv" else "application/zip"
